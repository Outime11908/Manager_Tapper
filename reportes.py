"""
Genera los reportes descargables (Excel y PDF) de Manager Tapper.
Cada función regresa un BytesIO listo para mandarse con send_file.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from fpdf import FPDF


def generar_excel_ventas(tickets) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    encabezados = ["Ticket", "Fecha", "Método", "Cliente", "Artículos", "Pagado", "Total"]
    ws.append(encabezados)
    fill = PatternFill(start_color="1E2B22", end_color="1E2B22", fill_type="solid")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FBF7EE")
        celda.fill = fill

    for t in tickets:
        ws.append([
            t.id,
            t.fecha.strftime("%d/%m/%Y %H:%M"),
            t.metodo_pago.capitalize(),
            t.cliente.nombre if t.cliente else "",
            len(t.items),
            "Sí" if t.pagado else "No",
            t.total,
        ])

    for col, ancho in zip("ABCDEFG", [10, 18, 12, 20, 12, 10, 12]):
        ws.column_dimensions[col].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_pdf_resumen(resumen_data, tickets) -> io.BytesIO:
    pdf = FPDF()
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Manager Tapper - Reporte", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 6, f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Resumen financiero", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    filas_resumen = [
        ("Capital invertido", resumen_data["capital_invertido"]),
        ("Ganancia bruta", resumen_data["ganancia_bruta"]),
        ("Gastos operativos", resumen_data["gastos_totales"]),
        ("Ganancia neta", resumen_data["ganancia_neta"]),
    ]
    for etiqueta, valor in filas_resumen:
        pdf.cell(60, 6, etiqueta)
        pdf.cell(0, 6, f"${valor:,.2f}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Historial de ventas", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)

    encabezados = ["Fecha", "Metodo", "Cliente", "Articulos", "Total"]
    anchos = (38, 28, 55, 25, 25)
    alineacion = ("LEFT", "LEFT", "LEFT", "RIGHT", "RIGHT")

    if not tickets:
        pdf.cell(0, 6, "Sin ventas registradas todavia.", new_x="LMARGIN", new_y="NEXT")
    else:
        with pdf.table(col_widths=anchos, text_align=alineacion) as table:
            fila = table.row()
            for h in encabezados:
                fila.cell(h)
            for t in tickets:
                fila = table.row()
                fila.cell(t.fecha.strftime("%d/%m/%Y %H:%M"))
                fila.cell(t.metodo_pago.capitalize())
                fila.cell(t.cliente.nombre if t.cliente else "-")
                fila.cell(str(len(t.items)))
                fila.cell(f"${t.total:,.2f}")

    buffer = io.BytesIO(bytes(pdf.output()))
    buffer.seek(0)
    return buffer
