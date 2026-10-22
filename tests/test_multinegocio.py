"""
Prueba de aislamiento entre negocios (multi-tenant).

Uso (desde la carpeta del proyecto):
    python tests/test_multinegocio.py

Crea una base de datos temporal, registra varios negocios y verifica que un
negocio NO puede ver, editar, borrar ni usar datos de otro. Si algo falla,
imprime qué y termina con error.
"""
import io
import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
_tmp.close()
os.environ["MANAGER_TAPPER_DB"] = _tmp.name  # antes de importar app

from app import app  # noqa: E402
from models import db, Negocio, Producto, Usuario  # noqa: E402

app.config["TESTING"] = True
fallos = []
total = 0


def check(descripcion, condicion):
    global total
    total += 1
    if not condicion:
        fallos.append(descripcion)
        print(f"  FALLA: {descripcion}")


def csrf_de(html):
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    return m.group(1) if m else ""


def registrar(client, negocio, username, email, password="clavesegura1", acepto=True, token=None):
    if token is None:
        token = csrf_de(client.get("/registro").get_data(as_text=True))
    data = {"csrf_token": token, "negocio": negocio, "username": username,
            "email": email, "password": password}
    if acepto:
        data["acepto_terminos"] = "on"
    return client.post("/registro", data=data)


def j(resp):
    return resp.get_json()


print("== 1. validaciones del registro ==")
c = app.test_client()
r = registrar(c, "X", "x1", "x1@test.com", acepto=False)
check("sin aceptar términos se rechaza", b"aceptar los t" in r.data and r.status_code == 200)
r = registrar(c, "X", "x1", "x1@test.com", password="corta")
check("contraseña corta se rechaza", b"al menos 8" in r.data)
r = registrar(c, "X", "x1", "correo-malo")
check("correo inválido se rechaza", b"correo v" in r.data)
r = registrar(c, "X", "x1", "x1@test.com", token="token-falso")
check("token CSRF falso se rechaza", "sesi".encode() in r.data and b"expir" in r.data)
with app.app_context():
    check("ninguno de esos intentos fallidos creó cuentas", Usuario.query.count() == 0)

print("== 2. registro de dos negocios ==")
A = app.test_client()
B = app.test_client()
rA = registrar(A, "Tienda A", "dueñoA", "a@test.com")
check("A se registra y entra (302 a /)", rA.status_code == 302 and rA.headers["Location"].endswith("/"))
rB = registrar(B, "Tienda B", "dueñoB", "b@test.com")
check("B se registra y entra", rB.status_code == 302)
D = app.test_client()
r = registrar(D, "Otra", "dueñoA", "otro@test.com")
check("usuario repetido se rechaza", b"ya est" in r.data)
r = registrar(D, "Otra", "otro", "a@test.com")
check("correo repetido se rechaza", b"ya existe" in r.data.lower())
check("/setup redirige a /registro", app.test_client().get("/setup").headers["Location"].endswith("/registro"))
check("con sesión, /registro manda a /", A.get("/registro").status_code == 302)

print("== 3. A crea sus datos ==")
cat_a = j(A.post("/api/categorias", json={"nombre": "Ropa"}))
prod_a = j(A.post("/api/productos", json={"nombre": "Camiseta", "costo": 50, "precio": 100, "stock": 20, "categoria_id": cat_a["id"]}))
cli_a = j(A.post("/api/clientes", json={"nombre": "Juan"}))
A.post("/api/capital", json={"monto": 1000})
A.post("/api/gastos", json={"descripcion": "Renta", "monto": 200})
A.post("/api/metas", json={"descripcion": "Meta", "monto_objetivo": 500, "fecha_inicio": "2020-01-01", "fecha_fin": "2099-12-31"})
t1 = j(A.post("/api/ventas", json={"items": [{"producto_id": prod_a["id"], "cantidad": 2}], "metodo_pago": "efectivo", "monto_pagado": 500}))
t2 = j(A.post("/api/ventas", json={"items": [{"producto_id": prod_a["id"], "cantidad": 1}], "metodo_pago": "credito", "cliente_id": cli_a["id"]}))
check("A tiene sus datos creados", cat_a.get("id") and prod_a.get("id") and cli_a.get("id") and t1.get("id") and t2.get("id"))
check("A ve su producto con stock 17", j(A.get("/api/productos"))[0]["stock"] == 17)

print("== 4. B puede usar los mismos nombres ==")
r = B.post("/api/categorias", json={"nombre": "Ropa"})
check("B crea categoría 'Ropa' igual que A", r.status_code == 201)
cat_b = j(r)
r = B.post("/api/productos", json={"nombre": "Camiseta", "costo": 10, "precio": 30, "stock": 5, "categoria_id": cat_b["id"]})
check("B crea producto 'Camiseta' igual que A", r.status_code == 201)
prod_b = j(r)
check("duplicado DENTRO de B se rechaza (400, no 500)", B.post("/api/productos", json={"nombre": "Camiseta", "costo": 1, "precio": 2, "stock": 1}).status_code == 400)
check("categoría duplicada DENTRO de B se rechaza", B.post("/api/categorias", json={"nombre": "Ropa"}).status_code == 400)

print("== 5. B no ve nada de A en los listados ==")
fresh = app.test_client()
registrar(fresh, "Tienda C", "dueñoC", "c@test.com")
C = fresh
for ruta in ["/api/clientes", "/api/gastos", "/api/capital", "/api/tickets", "/api/cuentas-por-cobrar", "/api/metas", "/api/ranking", "/api/inventario/alertas", "/api/categorias", "/api/productos"]:
    check(f"C (sin datos) ve vacío {ruta}", j(C.get(ruta)) == [])
res_c = j(C.get("/api/resumen"))
check("resumen de C en ceros", res_c["capital_invertido"] == 0 and res_c["ganancia_bruta"] == 0 and res_c["gastos_totales"] == 0)
check("ventas por periodo de C vacío", j(C.get("/api/ventas-por-periodo?agrupacion=dia")) == [])
check("B solo ve SU producto", [p["id"] for p in j(B.get("/api/productos"))] == [prod_b["id"]])
check("B solo ve SU categoría", [x["id"] for x in j(B.get("/api/categorias"))] == [cat_b["id"]])
check("B no ve el cliente de A", j(B.get("/api/clientes")) == [])
check("B no ve tickets de A", j(B.get("/api/tickets")) == [])
check("B no ve cuentas por cobrar de A", j(B.get("/api/cuentas-por-cobrar")) == [])
res_b = j(B.get("/api/resumen"))
check("resumen de B sin dinero de A", res_b["capital_invertido"] == 0 and res_b["ganancia_bruta"] == 0 and res_b["gastos_totales"] == 0)
res_a = j(A.get("/api/resumen"))
check("resumen de A correcto (cap 1000, gan 150, gastos 200)", res_a["capital_invertido"] == 1000 and res_a["ganancia_bruta"] == 150 and res_a["gastos_totales"] == 200)

print("== 6. B no puede tocar objetos de A por id ==")
pid, tid = prod_a["id"], t2["id"]
check("GET producto historial de A -> 404", B.get(f"/api/productos/{pid}/historial").status_code == 404)
check("PUT producto de A -> 404", B.put(f"/api/productos/{pid}", json={"nombre": "hackeado"}).status_code == 404)
check("PUT precio de A -> 404", B.put(f"/api/productos/{pid}/precio", json={"precio": 1}).status_code == 404)
check("DELETE producto de A -> 404", B.delete(f"/api/productos/{pid}").status_code == 404)
check("GET ticket de A -> 404", B.get(f"/api/tickets/{tid}").status_code == 404)
check("marcar-pagado de ticket de A -> 404", B.post(f"/api/tickets/{tid}/marcar-pagado").status_code == 404)
check("conteo de producto de A -> 404", B.post("/api/inventario/conteo", json={"producto_id": pid, "stock_contado": 0}).status_code == 404)
check("simular precio de producto de A -> 404", B.post("/api/simular-precio", json={"producto_id": pid, "porcentaje": 10}).status_code == 404)
check("vender producto de A -> 400", B.post("/api/ventas", json={"items": [{"producto_id": pid, "cantidad": 1}], "metodo_pago": "tarjeta"}).status_code == 400)
check("venta a crédito con cliente de A -> 400", B.post("/api/ventas", json={"items": [{"producto_id": prod_b["id"], "cantidad": 1}], "metodo_pago": "credito", "cliente_id": cli_a["id"]}).status_code == 400)
check("crear producto con categoría de A -> 400", B.post("/api/productos", json={"nombre": "Nuevo", "costo": 1, "precio": 2, "stock": 1, "categoria_id": cat_a["id"]}).status_code == 400)
check("editar producto propio con categoría de A -> 400", B.put(f"/api/productos/{prod_b['id']}", json={"categoria_id": cat_a["id"]}).status_code == 400)
check("producto_id basura -> 400 (no 500)", B.post("/api/simular-precio", json={"producto_id": "abc", "porcentaje": 1}).status_code == 400)

print("== 7. los datos de A siguen intactos ==")
pa = [p for p in j(A.get("/api/productos")) if p["id"] == pid][0]
check("producto de A sin cambios (nombre, stock 17, precio 100)", pa["nombre"] == "Camiseta" and pa["stock"] == 17 and pa["precio_actual"] == 100)
check("ticket de A sigue pendiente de pago", [t for t in j(A.get("/api/cuentas-por-cobrar"))][0]["id"] == tid)

print("== 8. reportes por negocio ==")
from openpyxl import load_workbook  # noqa: E402
rx = C.get("/api/reportes/ventas.xlsx")
wb = load_workbook(io.BytesIO(rx.data))
check("Excel de C sin ventas (solo encabezado)", wb.active.max_row == 1)
wa = load_workbook(io.BytesIO(A.get("/api/reportes/ventas.xlsx").data))
check("Excel de A con sus 2 ventas", wa.active.max_row == 3)
check("PDF de A se genera", A.get("/api/reportes/resumen.pdf").data[:4] == b"%PDF")
check("PDF de C se genera", C.get("/api/reportes/resumen.pdf").data[:4] == b"%PDF")
with app.app_context():
    n = db.session.get(Negocio, 3)
    n.nombre = "Tienda 🚀 €"
    db.session.commit()
check("PDF con nombre de negocio con emoji sigue funcionando", C.get("/api/reportes/resumen.pdf").data[:4] == b"%PDF")

print("== 9. sesión manipulada / sin sesión ==")
check("API sin sesión -> 401", app.test_client().get("/api/productos").status_code == 401)
check("página sin sesión -> redirige a login", app.test_client().get("/").status_code == 302)
M = app.test_client()
with M.session_transaction() as s:
    with app.app_context():
        uA = Usuario.query.filter_by(username="dueñoA").first()
        nB = Negocio.query.filter_by(nombre="Tienda B").first()
        s["usuario_id"], s["negocio_id"] = uA.id, nB.id   # usuario de A con negocio de B
check("sesión con usuario de A y negocio de B -> rechazada", M.get("/api/productos").status_code == 401)

print("== 10. suspender un negocio ==")
with app.app_context():
    nA = Negocio.query.filter_by(nombre="Tienda A").first()
    nA.activo = False
    db.session.commit()
check("A suspendido: su sesión deja de servir", A.get("/api/productos").status_code == 401)
L = app.test_client()
tok = csrf_de(L.get("/login").get_data(as_text=True))
r = L.post("/login", data={"csrf_token": tok, "username": "dueñoA", "password": "clavesegura1"})
check("A suspendido no puede iniciar sesión", b"suspendida" in r.data)
check("B sigue funcionando normal", B.get("/api/productos").status_code == 200)
with app.app_context():
    Negocio.query.filter_by(nombre="Tienda A").first().activo = True
    db.session.commit()
tok = csrf_de(L.get("/login").get_data(as_text=True))
r = L.post("/login", data={"csrf_token": tok, "username": "dueñoA", "password": "clavesegura1"})
check("reactivado, A inicia sesión (302)", r.status_code == 302)
check("y ve sus datos de siempre", len(j(L.get("/api/tickets"))) == 2)
L.get("/logout")
check("tras logout la API da 401", L.get("/api/tickets").status_code == 401)

print("== 11. límite de registros por conexión ==")
creadas = 3  # A, B y C ya cuentan; falta llegar a 5
for i in range(5 - creadas):
    r = registrar(app.test_client(), f"N{i}", f"u{i}", f"u{i}@test.com")
    check(f"registro extra {i} permitido", r.status_code == 302)
r = registrar(app.test_client(), "Exceso", "uexceso", "exceso@test.com")
check("el registro #6 desde la misma conexión se bloquea", "demasiadas".encode() in r.data.lower())

print()
if fallos:
    print(f"RESULTADO: {len(fallos)} FALLA(S) de {total} comprobaciones")
    for f in fallos:
        print("  -", f)
    sys.exit(1)
print(f"RESULTADO: {total} comprobaciones OK")
os.unlink(_tmp.name)
